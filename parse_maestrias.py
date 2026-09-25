import openpyxl, re
wb = openpyxl.load_workbook('uploads/FORMATOS_PNFA_MAESTRIA_CEUNEM_REVISADOS_(1)_(1).xlsx', data_only=True)

SHEET_TO_OFICIAL = {
 'MAESTRIA EDUCACION INICIAL': 'MAESTRÍA EN EDUCACIÓN INICIAL',
 'MAESTRIA EDUCACION PRIMARIA': 'MAESTRÍA EN EDUCACIÓN PRIMARIA',
 'MAESTRIA CIENCIAS': 'MAESTRÍA EN CIENCIAS NATURALES PARA EDUCACIÓN MEDIA',
 'MAESTRIA MATEMATICA': 'MAESTRÍA EN MATEMÁTICAS PARA EDUCACIÓN MEDIA',
 'MAESTRIA GHC': 'MAESTRÍA EN GEOGRAFÍA, HISTORIA Y CIUDADANÍA PARA EDUCACIÓN MEDIA',
 'MAESTRIA ING PARA EDUC MED': 'MAESTRÍA EN INGLÉS PARA EDUCACIÓN MEDIA',
 'MAESTRIA EDUCACIÓN FISICA': 'MAESTRÍA EN EDUCACIÓN FÍSICA PARA EDUCACIÓN MEDIA',
 'MAESTRIA DYSE': 'MAESTRÍA EN DIRECCIÓN Y SUPERVISIÓN EDUCATIVA',
 'MAESTRIA CULTURA': 'MAGÍSTER EN PEDAGOGÍA CULTURAL E INTERCULTURALIDAD',
 'MAESTRIA EDUC INDIGENA': 'MAGÍSTER EN EDUCACIÓN INDÍGENA',
 'MAESTRIA FRONTERA': 'MAGÍSTER EN EDUCACIÓN EN FRONTERA',
 'MAESTRIA AFRO': 'MAGÍSTER EN EDUCACIÓN Y PEDAGOGÍAS AFROVENEZOLANAS',
 'MAESTRIA LENG Y COM': 'MAESTRÍA EN LENGUA Y COMUNICACIÓN PARA EDUCACIÓN MEDIA',
}
TRAY_NAME = {1:'PRIMER TRAYECTO',2:'SEGUNDO TRAYECTO',3:'TERCER TRAYECTO',4:'CUARTO TRAYECTO'}

def clean(s):
    s = re.sub(r'\s+', ' ', str(s).replace('\n',' ')).strip()
    return s

def find_header_row(ws):
    for r in range(1, ws.max_row+1):
        vals = [clean(ws.cell(r,c).value or '') for c in range(1, ws.max_column+1)]
        if 'CEDULA' in [v.upper() for v in vals] and 'UC' in [v.upper() for v in vals]:
            return r
    return None

result = {}
for sh, oficial in SHEET_TO_OFICIAL.items():
    ws = wb[sh]
    hr = find_header_row(ws)
    tray = 0
    materias = []  # (tray, materia)
    for c in range(1, ws.max_column+1):
        v = clean(ws.cell(hr,c).value or '')
        up = v.upper()
        if up == 'UC':
            tray += 1
            continue
        if up == 'PERIODO' or up == '' or up.startswith('CALIFICAC'):
            continue
        # columnas antes del primer UC (datos del estudiante) -> ignorar
        if tray == 0:
            continue
        if up in ('UNIDAD CURRICULAR',):   # placeholder -> pendiente
            continue
        if tray <= 4:
            materias.append((tray, v))
    result[oficial] = materias

# Imprime resumen y el literal
for oficial, mats in result.items():
    print('###', oficial, '| total materias:', len(mats))
    for t,m in mats:
        print('   T%d | %s' % (t, m))
    print()
